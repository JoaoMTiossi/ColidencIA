"""
Análise IA v4 — raciocínio de advogado marcário sobre os 3.478 casos do engine v4.
Avalia: score ponderado (distintividade), score bruto (conjunto), spec overlap.
Gera Excel final colorido + resumo executivo.
"""
from __future__ import annotations
import sys, os, re
sys.path.insert(0, os.path.dirname(__file__))

import pandas as pd
from colisao_rpi.engine.similarity import similarity_score, weighted_similarity
from colisao_rpi.engine.nucleus import extract_nucleus
from colisao_rpi.engine.normalize import phonetic_key
from colisao_rpi.engine.rules import _spec_overlap, _THRESHOLD_PONDERADO_MESMA, _THRESHOLD_PONDERADO_CORRELATA, _CONJUNTO_OVERRIDE
from colisao_rpi.engine.corpus import load_corpus
from colisao_rpi.data.nice_matrix import classes_collide
from colisao_rpi.data.loader_rpi import load_rpi_records

_corpus_path = os.path.join(os.path.dirname(__file__), 'colisao_rpi', 'data', 'corpus_freq.json')
corpus = load_corpus(_corpus_path)

def get_cl(s):
    m = re.search(r'NCL\(13\)\s*(\d+)', str(s))
    return int(m.group(1)) if m else 0

def score_num(v):
    try: return float(v)
    except: return -1.0

def _cl_label(cl):
    L = {35:'serviços comerciais',36:'financeiro/seguros',41:'educação/entret.',
         43:'alimentação',44:'saúde/veterinária',42:'TI/científico',
         37:'construção',39:'transporte',30:'alimentos/confeit.',38:'telecom',
         40:'tratamento mat.',25:'vestuário',9:'eletrônicos',16:'papelaria',
         20:'móveis',29:'alimentos proc.',32:'bebidas n-alc.',33:'bebidas alc.'}
    return L.get(cl, f'cl.{cl}')

def veredito_ia(cli, rpi, cl_cli, cl_rpi_list, spec_cli, spec_rpi):
    cl_rpi    = cl_rpi_list[0] if cl_rpi_list else 0
    mesma_cl  = cl_cli == cl_rpi
    correlata = any(classes_collide(cl_cli, c) for c in cl_rpi_list) and not mesma_cl

    sc_bruto  = similarity_score(cli, rpi)
    sc_pond   = weighted_similarity(cli, rpi, cl_cli, corpus)
    sc_nuc    = similarity_score(extract_nucleus(cli), extract_nucleus(rpi))
    sc_spec   = _spec_overlap(spec_cli, spec_rpi)
    identico  = phonetic_key(cli) == phonetic_key(rpi)
    nuc_c     = extract_nucleus(cli)
    nuc_r     = extract_nucleus(rpi)

    threshold  = _THRESHOLD_PONDERADO_MESMA if mesma_cl else _THRESHOLD_PONDERADO_CORRELATA

    # --- Spec ---
    if sc_spec < 0:
        spec_txt = 'spec não disponível'
        spec_ok  = True   # sem info → não bloqueia
    elif sc_spec >= 0.35:
        spec_txt = f'specs compatíveis ({sc_spec:.0%})'
        spec_ok  = True
    elif sc_spec >= 0.08:
        spec_txt = f'specs parcialmente compatíveis ({sc_spec:.0%})'
        spec_ok  = True
    else:
        spec_txt = f'specs incompatíveis ({sc_spec:.0%})'
        spec_ok  = False

    cl_ctx = f'mesma classe {cl_cli}' if mesma_cl else f'classes correlatas {cl_cli}/{cl_rpi}'

    # --- Veredito ---
    if identico and spec_ok:
        v, p = 'ALTO_RISCO', 1
        motivo = (f'Marcas foneticamente idênticas ("{cli}" ≡ "{rpi}"), {cl_ctx}. '
                  f'{spec_txt}. Impedimento de registro — art. 124-XIX LPI.')
    elif identico and not spec_ok:
        v, p = 'MEDIO_RISCO', 2
        motivo = (f'Marcas idênticas ("{cli}" ≡ "{rpi}") mas {spec_txt}. '
                  f'Risco de associação de origem — advogado deve confirmar.')
    elif sc_pond >= threshold and spec_ok and mesma_cl:
        v, p = 'ALTO_RISCO', 1
        motivo = (f'Score ponderado {sc_pond:.2f} ≥ {threshold} na {cl_ctx}. '
                  f'Núcleos: "{nuc_c}" ≈ "{nuc_r}". {spec_txt}. '
                  f'Elementos distintivos similares — risco concreto de confusão.')
    elif sc_pond >= threshold and spec_ok and correlata:
        v, p = 'MEDIO_RISCO', 2
        motivo = (f'Score ponderado {sc_pond:.2f} ≥ {threshold} em {cl_ctx}. '
                  f'Núcleos: "{nuc_c}" ≈ "{nuc_r}". {spec_txt}. '
                  f'Risco de associação indireta.')
    elif sc_pond >= threshold and not spec_ok:
        v, p = 'BAIXO_RISCO', 3
        motivo = (f'Score ponderado {sc_pond:.2f} mas {spec_txt}. '
                  f'Elementos distintivos similares mas segmentos distintos — monitorar.')
    elif sc_bruto >= _CONJUNTO_OVERRIDE and spec_ok:
        v, p = 'ALTO_RISCO', 1
        motivo = (f'Conjunto marcário quasi-idêntico (score bruto {sc_bruto:.2f} ≥ {_CONJUNTO_OVERRIDE}). '
                  f'{spec_txt}, {cl_ctx}. Consumidor médio pode confundir.')
    elif sc_bruto >= _CONJUNTO_OVERRIDE and not spec_ok:
        v, p = 'MEDIO_RISCO', 2
        motivo = (f'Conjunto quasi-idêntico (bruto {sc_bruto:.2f}) mas {spec_txt}. Revisar.')
    elif sc_nuc >= 0.90 and mesma_cl and spec_ok:
        v, p = 'ALTO_RISCO', 1
        motivo = (f'Núcleos quasi-idênticos "{nuc_c}" ≈ "{nuc_r}" (score {sc_nuc:.2f}), '
                  f'{cl_ctx}. {spec_txt}.')
    elif sc_nuc >= 0.80 and mesma_cl and spec_ok:
        v, p = 'MEDIO_RISCO', 2
        motivo = (f'Núcleos similares "{nuc_c}" ≈ "{nuc_r}" (score {sc_nuc:.2f}), '
                  f'{cl_ctx}. {spec_txt}. Avaliar apresentação visual.')
    elif sc_pond < 0.45 or not spec_ok:
        v, p = 'FALSO_POSITIVO', 4
        motivo = (f'Score ponderado baixo ({sc_pond:.2f}) e/ou {spec_txt}. '
                  f'Sem elemento distintivo em comum — não há risco real.')
    else:
        v, p = 'BAIXO_RISCO', 3
        motivo = (f'Score ponderado {sc_pond:.2f}, bruto {sc_bruto:.2f}, {cl_ctx}. '
                  f'{spec_txt}. Monitoramento passivo.')

    return {'VEREDITO': v, 'PRIORIDADE': p,
            'SC_POND': round(sc_pond,3), 'SC_BRUTO': round(sc_bruto,3),
            'SC_NUC': round(sc_nuc,3),
            'SC_SPEC': round(sc_spec,3) if sc_spec >= 0 else 'N/D',
            'NUCLEO_CLI': nuc_c, 'NUCLEO_RPI': nuc_r,
            'MOTIVO_IA': motivo}

# ---------------------------------------------------------------------------
print('Carregando specs...')
XLSX = r'C:\Users\jolut\OneDrive\Desktop\COLIDENCIA\zipppp\RM2886\Relatório_Detalhado_Marcas -18-05-2026 - Total de processos - 46.619.xlsx'
XML  = r'C:\Users\jolut\OneDrive\Desktop\COLIDENCIA\zipppp\RM2886\RM2886.xml'

df_cli = pd.read_excel(XLSX, dtype=str)
df_cli.columns = [c.strip() for c in df_cli.columns]
spec_cli_map = dict(zip(df_cli['PROCESSO'].astype(str).str.strip(), df_cli['ESPECIFICAÇÃO'].fillna('')))

rpi_records, _, _ = load_rpi_records(XML)
spec_rpi_map = {r['processo']: r['especificacoes'] for r in rpi_records}

print('Carregando relatório v4...')
df = pd.read_excel('./relatorio_v4/Relatorio_Colidencia_RPI_2886_28-04-2026_A_PROVINCIA.xlsx', header=6)
df.columns = ['PROC_CLI','MARCA_CLI','CLASSE_CLI','TITULAR_CLI','PROC_3O','MARCA_3O','CLASSE_3O']
df = df[df['PROC_CLI'] != 'PROCESSO CLIENTE'].dropna(subset=['PROC_CLI'])
print(f'  {len(df)} casos')

rows = []
for i, (_, r) in enumerate(df.iterrows(), 1):
    if i % 500 == 0:
        print(f'  ... {i}/{len(df)}')
    cli    = str(r['MARCA_CLI'])
    rpi_n  = str(r['MARCA_3O'])
    cl_c   = get_cl(str(r['CLASSE_CLI']))
    cls_r  = [get_cl(x) for x in str(r['CLASSE_3O']).split(',')]
    proc_c = str(r['PROC_CLI']).strip()
    proc_r = str(r['PROC_3O']).strip()
    cl_r0  = str(cls_r[0]) if cls_r else ''

    spec_c = spec_cli_map.get(proc_c, '')
    spec_r_d = spec_rpi_map.get(proc_r, {})
    spec_r = spec_r_d.get(cl_r0, '') or next(iter(spec_r_d.values()), '')

    v = veredito_ia(cli, rpi_n, cl_c, cls_r, spec_c, spec_r)

    rows.append({'PROCESSO CLIENTE': r['PROC_CLI'], 'MARCA CLIENTE': cli,
                 'CLASSE CLIENTE': r['CLASSE_CLI'], 'TITULAR CLIENTE': r['TITULAR_CLI'],
                 'PROCESSO TERCEIRO': r['PROC_3O'], 'MARCA TERCEIRO': rpi_n,
                 'CLASSE TERCEIRO': r['CLASSE_3O'], **v,
                 'ESPEC_CLI': str(spec_c)[:200], 'ESPEC_RPI': str(spec_r)[:200]})

df_out = pd.DataFrame(rows).sort_values(['PRIORIDADE','SC_POND'], ascending=[True,False])

# ---------------------------------------------------------------------------
cnt = df_out['VEREDITO'].value_counts()
total = len(df_out)
print()
print('='*60)
print('RESULTADO FINAL — ANÁLISE IA v4')
print('='*60)
for v, n in cnt.items():
    print(f'  {v:<20} {n:>5} ({n/total*100:.1f}%)')
print(f'  {"TOTAL":<20} {total:>5}')

print()
print('TOP 15 ALTO_RISCO:')
for _, r in df_out[df_out['VEREDITO']=='ALTO_RISCO'].head(15).iterrows():
    print(f"  [{r['SC_POND']:.3f}/{r['SC_BRUTO']:.3f}] cl{r['CLASSE CLIENTE'][-2:]}  "
          f"{r['MARCA CLIENTE']!r:38} x {r['MARCA TERCEIRO']!r}")

# ---------------------------------------------------------------------------
OUTPUT = './relatorio_v4/Relatorio_Final_IA_v4_RPI2886.xlsx'
print(f'\nSalvando {OUTPUT}...')

cores = {'ALTO_RISCO':'FF9999','MEDIO_RISCO':'FFE599',
         'BAIXO_RISCO':'C6EFCE','FALSO_POSITIVO':'D9D9D9'}

with pd.ExcelWriter(OUTPUT, engine='openpyxl') as writer:
    # Aba completa
    df_out.to_excel(writer, index=False, sheet_name='Todos')
    # Aba alto risco
    df_out[df_out['VEREDITO']=='ALTO_RISCO'].to_excel(writer, index=False, sheet_name='Alto_Risco')
    # Aba médio risco
    df_out[df_out['VEREDITO']=='MEDIO_RISCO'].to_excel(writer, index=False, sheet_name='Medio_Risco')
    # Abas top 5 classes
    cl_counts = df_out[df_out['VEREDITO'].isin(['ALTO_RISCO','MEDIO_RISCO'])].groupby(
        df_out['CLASSE CLIENTE'].str.extract(r'(\d+)$')[0]).size().sort_values(ascending=False)
    for cl in cl_counts.head(5).index:
        sub = df_out[df_out['CLASSE CLIENTE'].str.contains(f'NCL\\(13\\) {cl}\\b', regex=True)]
        sub = sub[sub['VEREDITO'].isin(['ALTO_RISCO','MEDIO_RISCO'])]
        if len(sub):
            nome = f'Cl{cl}_{_cl_label(int(cl))[:8]}'
            nome = re.sub(r'[/\\?*\[\]:]', '', nome)[:31]
            sub.to_excel(writer, index=False, sheet_name=nome)

    from openpyxl.styles import PatternFill, Font, Alignment
    widths = {'PROCESSO CLIENTE':15,'MARCA CLIENTE':36,'CLASSE CLIENTE':13,
              'TITULAR CLIENTE':34,'PROCESSO TERCEIRO':15,'MARCA TERCEIRO':36,
              'CLASSE TERCEIRO':13,'VEREDITO':15,'PRIORIDADE':10,
              'SC_POND':10,'SC_BRUTO':10,'SC_NUC':10,'SC_SPEC':10,
              'NUCLEO_CLI':22,'NUCLEO_RPI':22,'MOTIVO_IA':80,
              'ESPEC_CLI':44,'ESPEC_RPI':44}
    hf = PatternFill('solid', fgColor='1F3864')

    for sn in writer.sheets:
        ws = writer.sheets[sn]
        for cell in ws[1]:
            cell.fill = hf
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
        try:
            col_v = next(c.column for c in ws[1] if c.value == 'VEREDITO')
        except StopIteration:
            col_v = None
        for row in ws.iter_rows(min_row=2):
            v = row[col_v-1].value if col_v else ''
            fill = PatternFill('solid', fgColor=cores.get(v or '', 'FFFFFF'))
            for cell in row:
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            if col_v:
                row[col_v-1].font = Font(bold=True)
        for col in ws.iter_cols(1, ws.max_column):
            h = col[0].value or ''
            ws.column_dimensions[col[0].column_letter].width = widths.get(h, 13)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

print('Concluído.')
