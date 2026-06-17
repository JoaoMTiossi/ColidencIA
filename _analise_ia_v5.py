"""
Análise IA v5 — raciocínio explícito de distintividade.

A IA entende quais tokens são descritores e quais são distintivos, e fundamenta
o parecer nisso: "ESFIHA é descritor declarado da cl.30 (42/412 specs) — a
comparação real é DR. vs Go — distintos — não colide."
"""
from __future__ import annotations
import sys, os, re
sys.path.insert(0, os.path.dirname(__file__))

import pandas as pd
from colisao_rpi.engine.similarity import similarity_score, weighted_similarity
from colisao_rpi.engine.nucleus import extract_nucleus
from colisao_rpi.engine.normalize import phonetic_key, normalize, apply_phonetic
from colisao_rpi.engine.rules import (_spec_overlap,
    _THRESHOLD_PONDERADO_MESMA, _THRESHOLD_PONDERADO_CORRELATA, _CONJUNTO_OVERRIDE)
from colisao_rpi.engine.distinctiveness import distinctive_tokens, is_descriptive
from colisao_rpi.engine.corpus import load_spec_corpus, update_spec_corpus
from colisao_rpi.data.nice_matrix import classes_collide
from colisao_rpi.data.loader_rpi import load_rpi_records
from rapidfuzz import fuzz as _fuzz

_SPEC_CORPUS_PATH = os.path.join(
    os.path.dirname(__file__), 'colisao_rpi', 'data', 'corpus_spec_freq.json'
)

# ---------------------------------------------------------------------------
# Carregar / construir spec corpus
# ---------------------------------------------------------------------------
def get_cl(s):
    m = re.search(r'NCL\(13\)\s*(\d+)', str(s))
    return int(m.group(1)) if m else 0

def score_num(v):
    try: return float(v)
    except: return -1.0

def _cl_label(cl):
    L = {35:'serv.comerciais', 36:'financeiro/seguros', 41:'educação/entret.',
         43:'alimentação', 44:'saúde/vet.', 42:'TI/científico',
         37:'construção', 39:'transporte', 30:'alimentos/confeit.',
         38:'telecom', 40:'tratamento mat.', 25:'vestuário',
         9:'eletrônicos', 16:'papelaria', 20:'móveis',
         29:'alimentos proc.', 32:'bebidas n-alc.', 33:'bebidas alc.'}
    return L.get(cl, f'cl.{cl}')

# ---------------------------------------------------------------------------
# Raciocínio de distintividade para a IA
# ---------------------------------------------------------------------------
def _analisar_tokens(marca: str, classe: int, spec_corpus: dict) -> dict:
    """
    Classifica cada token da marca como DISTINTIVO ou DESCRITOR,
    identifica o elemento principal comparável, e explica o raciocínio.
    """
    dtoks = distinctive_tokens(marca, classe, spec_corpus)
    distintos  = [(tok, w, mot) for tok, w, is_d, mot in dtoks if not is_d and w > 0.3]
    descritores = [(tok, w, mot) for tok, w, is_d, mot in dtoks if is_d]

    elem_principal = ' '.join(t for t, _, _ in distintos) if distintos else None
    return {
        'distintos':    distintos,
        'descritores':  descritores,
        'elem_principal': elem_principal,
        'so_descritores': len(distintos) == 0,
    }


def veredito_ia_v5(cli: str, rpi: str, cl_cli: int, cl_rpi_list: list[int],
                   spec_cli: str, spec_rpi: str, spec_corpus: dict) -> dict:
    """
    Emite veredito com raciocínio explícito sobre distintividade.
    A IA analisa token a token antes de concluir.
    """
    cl_rpi   = cl_rpi_list[0] if cl_rpi_list else 0
    mesma_cl = cl_cli == cl_rpi
    correlata = any(classes_collide(cl_cli, c) for c in cl_rpi_list) and not mesma_cl

    sc_bruto = similarity_score(cli, rpi)
    sc_pond  = weighted_similarity(cli, rpi, cl_cli, spec_corpus)
    sc_nuc   = similarity_score(extract_nucleus(cli), extract_nucleus(rpi))
    sc_spec  = _spec_overlap(spec_cli, spec_rpi)
    identico = phonetic_key(cli) == phonetic_key(rpi)

    # Analisar tokens de cada marca
    ana_c = _analisar_tokens(cli, cl_cli, spec_corpus)
    ana_r = _analisar_tokens(rpi, cl_rpi, spec_corpus)

    threshold = _THRESHOLD_PONDERADO_MESMA if mesma_cl else _THRESHOLD_PONDERADO_CORRELATA
    cl_ctx = f'mesma classe {cl_cli}' if mesma_cl else f'classes {cl_cli}/{cl_rpi}'

    # --- Spec ---
    if sc_spec < 0:
        spec_txt, spec_ok = 'spec não disponível', True
    elif sc_spec >= 0.35:
        spec_txt, spec_ok = f'specs compatíveis ({sc_spec:.0%})', True
    elif sc_spec >= 0.08:
        spec_txt, spec_ok = f'specs parcialmente compatíveis ({sc_spec:.0%})', True
    else:
        spec_txt, spec_ok = f'specs incompatíveis ({sc_spec:.0%})', False

    # --- Construir raciocínio passo a passo ---
    raciocinio = []

    # Passo 1: analisar cada token
    if ana_c['descritores']:
        nomes = ', '.join(f'"{t}"' for t, _, _ in ana_c['descritores'])
        mots  = ana_c['descritores'][0][2] if ana_c['descritores'] else ''
        raciocinio.append(f'Em "{cli}": {nomes} é descritor ({mots})')
    if ana_r['descritores']:
        nomes = ', '.join(f'"{t}"' for t, _, _ in ana_r['descritores'])
        mots  = ana_r['descritores'][0][2] if ana_r['descritores'] else ''
        raciocinio.append(f'Em "{rpi}": {nomes} é descritor ({mots})')

    # Passo 2: elementos comparáveis
    elem_c = ana_c['elem_principal'] or extract_nucleus(cli)
    elem_r = ana_r['elem_principal'] or extract_nucleus(rpi)
    sc_elem = similarity_score(elem_c, elem_r) if elem_c and elem_r else 0.0
    raciocinio.append(
        f'Elementos distintivos: "{elem_c}" vs "{elem_r}" → similaridade {sc_elem:.2f}'
    )

    # Passo 3: spec
    raciocinio.append(f'Especificação: {spec_txt}')

    # --- Decisão ---
    ambas_so_descritores = ana_c['so_descritores'] and ana_r['so_descritores']

    if identico and spec_ok:
        v, p = 'ALTO_RISCO', 1
        decisao = (f'Marcas foneticamente idênticas — mesmo que os tokens sejam '
                   f'descritivos, a identidade total cria risco de confusão de origem. '
                   f'{spec_txt}, {cl_ctx}.')

    elif identico and not spec_ok:
        v, p = 'MEDIO_RISCO', 2
        decisao = f'Marcas idênticas mas {spec_txt} — advogado deve avaliar risco de associação.'

    elif ambas_so_descritores:
        v, p = 'FALSO_POSITIVO', 4
        descr_c = ', '.join(f'"{t}"' for t, _, _ in ana_c['descritores'])
        descr_r = ', '.join(f'"{t}"' for t, _, _ in ana_r['descritores'])
        decisao = (f'AMBAS as marcas são compostas apenas por descritores: '
                   f'"{cli}" ({descr_c}) e "{rpi}" ({descr_r}). '
                   f'Nenhuma tem elemento distintivo — não configura colidência.')

    elif ana_c['so_descritores'] or ana_r['so_descritores']:
        v, p = 'FALSO_POSITIVO', 4
        quem = cli if ana_c['so_descritores'] else rpi
        decisao = (f'"{quem}" é composta apenas por descritores — não tem elemento '
                   f'distintivo para comparar. Sem risco de confusão.')

    elif sc_elem < 0.55 and sc_pond < threshold:
        v, p = 'FALSO_POSITIVO', 4
        decisao = (f'Elementos distintivos diferentes: "{elem_c}" vs "{elem_r}" '
                   f'(score {sc_elem:.2f}). A similaridade vem dos descritores '
                   f'compartilhados, que não podem ser monopolizados.')

    elif sc_pond >= threshold and spec_ok and mesma_cl:
        v, p = 'ALTO_RISCO', 1
        decisao = (f'Elementos distintivos similares: "{elem_c}" ≈ "{elem_r}" '
                   f'(score ponderado {sc_pond:.2f} ≥ {threshold}), {cl_ctx}. '
                   f'{spec_txt}. Risco concreto de confusão — art. 124-XIX LPI.')

    elif sc_pond >= threshold and spec_ok and correlata:
        v, p = 'MEDIO_RISCO', 2
        decisao = (f'Elementos distintivos similares: "{elem_c}" ≈ "{elem_r}" '
                   f'(score ponderado {sc_pond:.2f}), {cl_ctx}. '
                   f'{spec_txt}. Risco de associação indireta.')

    elif sc_pond >= threshold and not spec_ok:
        v, p = 'BAIXO_RISCO', 3
        decisao = (f'Elementos distintivos similares ({sc_pond:.2f}) mas {spec_txt}. '
                   f'Monitorar expansão de cobertura.')

    elif sc_bruto >= _CONJUNTO_OVERRIDE and spec_ok:
        v, p = 'ALTO_RISCO', 1
        decisao = (f'Conjunto marcário quasi-idêntico (score bruto {sc_bruto:.2f}), '
                   f'{cl_ctx}. {spec_txt}.')

    elif sc_elem >= 0.80 and mesma_cl and spec_ok:
        v, p = 'ALTO_RISCO', 1
        decisao = (f'Elementos distintivos muito similares: "{elem_c}" ≈ "{elem_r}" '
                   f'(score {sc_elem:.2f}), {cl_ctx}. {spec_txt}.')

    elif sc_elem >= 0.70 and mesma_cl and spec_ok:
        v, p = 'MEDIO_RISCO', 2
        decisao = (f'Elementos distintivos moderadamente similares: "{elem_c}" ≈ "{elem_r}" '
                   f'(score {sc_elem:.2f}), {cl_ctx}. {spec_txt}. Avaliar visualmente.')

    elif not spec_ok:
        v, p = 'FALSO_POSITIVO', 4
        decisao = f'Elementos distintos + {spec_txt}. Sem risco real.'

    else:
        v, p = 'BAIXO_RISCO', 3
        decisao = (f'Elementos distintivos pouco similares: "{elem_c}" vs "{elem_r}" '
                   f'(score {sc_elem:.2f}). Monitoramento passivo.')

    motivo_completo = ' | '.join(raciocinio) + ' → ' + decisao

    return {
        'VEREDITO':    v,
        'PRIORIDADE':  p,
        'ELEM_CLI':    elem_c,
        'ELEM_RPI':    elem_r,
        'SC_ELEM':     round(sc_elem, 3),
        'SC_POND':     round(sc_pond, 3),
        'SC_BRUTO':    round(sc_bruto, 3),
        'SC_SPEC':     round(sc_spec, 3) if sc_spec >= 0 else 'N/D',
        'DESCR_CLI':   ', '.join(t for t, _, _ in ana_c['descritores']),
        'DESCR_RPI':   ', '.join(t for t, _, _ in ana_r['descritores']),
        'MOTIVO_IA':   motivo_completo,
    }


# ---------------------------------------------------------------------------
print('Carregando spec corpus (ou construindo)...')
XML  = r'C:\Users\jolut\OneDrive\Desktop\COLIDENCIA\zipppp\RM2886\RM2886.xml'
XLSX = r'C:\Users\jolut\OneDrive\Desktop\COLIDENCIA\zipppp\RM2886\Relatório_Detalhado_Marcas -18-05-2026 - Total de processos - 46.619.xlsx'

rpi_records, _, _ = load_rpi_records(XML)
spec_corpus = update_spec_corpus(rpi_records, _SPEC_CORPUS_PATH)

N_total = sum(spec_corpus.get('_N', {}).values())
print(f'  Spec corpus: {N_total} registros com spec, '
      f'{len(spec_corpus)-1} classes')

# Validar detecção de descritores conhecidos
print('\n  Teste de descritores:')
testes = [('ESFIHA',30),('LAVA',37),('JATO',37),('CONTABILIDADE',35),
          ('CORRETORA',36),('SUSHI',43),('AIKA',35),('VELTI',35),('EDIPHARMA',35)]
for tok, cl in testes:
    is_d, mot = is_descriptive(tok, cl, spec_corpus)
    flag = '✓ DESCRITOR' if is_d else '  DISTINTIVO'
    print(f'    {tok:<18} cl.{cl}  {flag}  {mot[:60]}')

print()

# Specs dos clientes
df_cli = pd.read_excel(XLSX, dtype=str)
df_cli.columns = [c.strip() for c in df_cli.columns]
spec_cli_map = dict(zip(df_cli['PROCESSO'].astype(str).str.strip(),
                        df_cli['ESPECIFICAÇÃO'].fillna('')))
spec_rpi_map = {r['processo']: r['especificacoes'] for r in rpi_records}

# ---------------------------------------------------------------------------
print('Carregando relatório v4...')
df = pd.read_excel('./relatorio_v4/Relatorio_Colidencia_RPI_2886_28-04-2026_A_PROVINCIA.xlsx', header=6)
df.columns = ['PROC_CLI','MARCA_CLI','CLASSE_CLI','TITULAR_CLI','PROC_3O','MARCA_3O','CLASSE_3O']
df = df[df['PROC_CLI'] != 'PROCESSO CLIENTE'].dropna(subset=['PROC_CLI'])
print(f'  {len(df)} casos')

rows = []
for i, (_, r) in enumerate(df.iterrows(), 1):
    if i % 500 == 0:
        print(f'  ... {i}/{len(df)}')

    cli   = str(r['MARCA_CLI'])
    rpi_n = str(r['MARCA_3O'])
    cl_c  = get_cl(str(r['CLASSE_CLI']))
    cls_r = [get_cl(x) for x in str(r['CLASSE_3O']).split(',')]
    proc_c = str(r['PROC_CLI']).strip()
    proc_r = str(r['PROC_3O']).strip()
    cl_r0  = str(cls_r[0]) if cls_r else ''

    spec_c   = spec_cli_map.get(proc_c, '')
    spec_r_d = spec_rpi_map.get(proc_r, {})
    spec_r   = spec_r_d.get(cl_r0, '') or next(iter(spec_r_d.values()), '')

    v = veredito_ia_v5(cli, rpi_n, cl_c, cls_r, spec_c, spec_r, spec_corpus)
    rows.append({
        'PROCESSO CLIENTE': r['PROC_CLI'], 'MARCA CLIENTE': cli,
        'CLASSE CLIENTE': r['CLASSE_CLI'], 'TITULAR CLIENTE': r['TITULAR_CLI'],
        'PROCESSO TERCEIRO': r['PROC_3O'], 'MARCA TERCEIRO': rpi_n,
        'CLASSE TERCEIRO': r['CLASSE_3O'], **v,
        'ESPEC_CLI': str(spec_c)[:200], 'ESPEC_RPI': str(spec_r)[:200],
    })

df_out = pd.DataFrame(rows).sort_values(['PRIORIDADE','SC_POND'], ascending=[True,False])

# ---------------------------------------------------------------------------
cnt   = df_out['VEREDITO'].value_counts()
total = len(df_out)
print()
print('='*60)
print('RESULTADO — ANÁLISE IA v5 (distintividade explícita)')
print('='*60)
for v, n in cnt.items():
    print(f'  {v:<22} {n:>5} ({n/total*100:.1f}%)')
print(f'  {"TOTAL":<22} {total:>5}')

print()
print('TOP 15 ALTO_RISCO:')
for _, r in df_out[df_out['VEREDITO']=='ALTO_RISCO'].head(15).iterrows():
    print(f"  [{r['SC_POND']:.3f}] cl{r['CLASSE CLIENTE'][-2:]}  "
          f'"{r["ELEM_CLI"]}" ≈ "{r["ELEM_RPI"]}"  |  '
          f'{r["MARCA CLIENTE"]!r} x {r["MARCA TERCEIRO"]!r}')

print()
print('EXEMPLOS DE FALSOS POSITIVOS IDENTIFICADOS (descritores):')
fp = df_out[df_out['VEREDITO']=='FALSO_POSITIVO'].head(10)
for _, r in fp.iterrows():
    print(f"  {r['MARCA CLIENTE']!r} x {r['MARCA TERCEIRO']!r}")
    print(f"    descr CLI={r['DESCR_CLI'] or '—'}  descr RPI={r['DESCR_RPI'] or '—'}")
    motivo_curto = r['MOTIVO_IA'].split('→')[-1].strip()[:120]
    print(f"    → {motivo_curto}")
    print()

# ---------------------------------------------------------------------------
OUTPUT = './relatorio_v4/Relatorio_Final_IA_v5_RPI2886.xlsx'
print(f'Salvando {OUTPUT}...')

cores = {'ALTO_RISCO':'FF9999','MEDIO_RISCO':'FFE599',
         'BAIXO_RISCO':'C6EFCE','FALSO_POSITIVO':'D9D9D9'}

with pd.ExcelWriter(OUTPUT, engine='openpyxl') as writer:
    df_out.to_excel(writer, index=False, sheet_name='Todos')
    df_out[df_out['VEREDITO']=='ALTO_RISCO'].to_excel(
        writer, index=False, sheet_name='Alto_Risco')
    df_out[df_out['VEREDITO']=='MEDIO_RISCO'].to_excel(
        writer, index=False, sheet_name='Medio_Risco')

    cl_top = df_out[df_out['VEREDITO'].isin(['ALTO_RISCO','MEDIO_RISCO'])]\
        .groupby(df_out['CLASSE CLIENTE'].str.extract(r'(\d+)$')[0])\
        .size().sort_values(ascending=False)
    for cl in cl_top.head(5).index:
        sub = df_out[
            df_out['CLASSE CLIENTE'].str.contains(f'NCL\\(13\\) {cl}\\b', regex=True) &
            df_out['VEREDITO'].isin(['ALTO_RISCO','MEDIO_RISCO'])
        ]
        if len(sub):
            nome = re.sub(r'[/\\?*\[\]:]', '', f'Cl{cl}_{_cl_label(int(cl))[:8]}')[:31]
            sub.to_excel(writer, index=False, sheet_name=nome)

    from openpyxl.styles import PatternFill, Font, Alignment
    widths = {
        'PROCESSO CLIENTE':15,'MARCA CLIENTE':34,'CLASSE CLIENTE':13,
        'TITULAR CLIENTE':32,'PROCESSO TERCEIRO':15,'MARCA TERCEIRO':34,
        'CLASSE TERCEIRO':13,'VEREDITO':15,'PRIORIDADE':10,
        'ELEM_CLI':20,'ELEM_RPI':20,'SC_ELEM':9,'SC_POND':9,
        'SC_BRUTO':9,'SC_SPEC':9,'DESCR_CLI':22,'DESCR_RPI':22,
        'MOTIVO_IA':90,'ESPEC_CLI':42,'ESPEC_RPI':42,
    }
    hf = PatternFill('solid', fgColor='1F3864')
    for sn in writer.sheets:
        ws = writer.sheets[sn]
        for cell in ws[1]:
            cell.fill = hf
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
        try:
            col_v = next(c.column for c in ws[1] if c.value == 'VEREDITO')
        except StopIteration:
            col_v = None
        for row in ws.iter_rows(min_row=2):
            v = row[col_v-1].value if col_v else ''
            for cell in row:
                cell.fill = PatternFill('solid', fgColor=cores.get(v or '', 'FFFFFF'))
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            if col_v:
                row[col_v-1].font = Font(bold=True)
        for col in ws.iter_cols(1, ws.max_column):
            ws.column_dimensions[col[0].column_letter].width = widths.get(col[0].value or '', 13)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

print('Concluído.')
