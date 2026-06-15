"""
Extrai tokens frequentes das marcas no relatório v2 e gera Excel para revisão.
Para cada token indica quantas marcas distintas o contêm e exemplos.
"""
from __future__ import annotations
import sys, os, re
sys.path.insert(0, os.path.dirname(__file__))

import pandas as pd
from collections import defaultdict
from colisao_rpi.engine.normalize import normalize
from colisao_rpi.engine.nucleus import ALL_STOPWORDS

# Já estão no filtro atual — não precisam aparecer na lista de revisão
JA_FILTRADOS: frozenset[str] = frozenset({
    'MOTORS','MOTOR','MOTO','AUTO','AUTOMOVEL','VEICULOS','VEICULO','AUTOMOTIVO',
    'DELIVERY','ENTREGA','EXPRESS','EXPRESSO',
    'SEGUROS','SEGURO','CORRETORA','FINANCEIRA','CREDITO','INVESTIMENTOS',
    'MODA','MODAS','FASHION','ROUPAS','VESTUARIO',
    'PIZZA','PIZZARIA','BURGER','FOOD','LANCHE','SUSHI','RESTAURANTE',
    'SHOP','STORE','MERCADO','MARKET','COMERCIO','LOJA',
    'TECH','DIGITAL','ONLINE','SISTEMAS','SOLUCOES',
    'TUDO','GERAL','CLEAN','TOTAL',
    'PINK','ROSA','VERDE','AZUL','BRANCO','PRETO','DOURADO','PRATA','OURO','BRONZE',
    'CASA','NOVA','NOVO',
    'ENGENHARIA','CONSTRUTORA','CONSTRUCAO','ARQUITETURA',
    'INFORMATICA','TECNOLOGIA','CLINICA','HOSPITAL','SAUDE',
    'ACADEMIA','ESCOLA','INSTITUTO','EDUCACAO',
    'ADVOCACIA','JURIDICO','ASSESSORIA','CONSULTORIA',
    # stopwords já existentes
    *ALL_STOPWORDS,
})

# Termos que o usuário já identificou como comuns (pré-marcar no Excel)
COMUNS_USUARIO: frozenset[str] = frozenset({
    'BAR', 'CAFE', 'RESTAURANTE', 'TOCA', 'POSTO', 'CENTRO',
    # variantes normalizadas
    'CAFETERIA', 'LANCHONETE', 'PADARIA', 'DOCERIA', 'PIZZARIA',
    'SUPERMERCADO', 'FARMACIA', 'DROGARIA',
    'ATELIE', 'ESTUDIO', 'STUDIO',
    'TURISMO', 'VIAGENS', 'TRANSPORTE',
    'SOLAR', 'PALACE', 'PLAZA', 'PARK',
    'MASTER', 'PRIME', 'GOLD', 'SILVER',
    'BRASIL', 'NACIONAL', 'REGIONAL',
    'NORTE', 'SUL', 'LESTE', 'OESTE', 'SUDESTE',
})

MIN_LEN = 3
MIN_FREQ = 5   # mínimo de marcas distintas para aparecer na lista

# ---------------------------------------------------------------------------
print('Carregando relatório v2...')
df = pd.read_excel(
    './relatorio_v2/Relatorio_Colidencia_RPI_2886_28-04-2026_A_PROVINCIA.xlsx',
    header=6
)
df.columns = ['PROCESSO_CLI','MARCA_CLI','CLASSE_CLI','TITULAR_CLI',
              'PROCESSO_3O','MARCA_3O','CLASSE_3O']
df = df[df['PROCESSO_CLI'] != 'PROCESSO CLIENTE'].dropna(subset=['PROCESSO_CLI'])
print(f'  {len(df)} pares de colidência')

# Coletar todas as marcas únicas (cliente + terceiro)
todas_marcas: set[str] = set(df['MARCA_CLI'].astype(str)) | set(df['MARCA_3O'].astype(str))
print(f'  {len(todas_marcas)} marcas únicas')

# Para cada token: lista de marcas que o contêm
token_marcas: dict[str, set[str]] = defaultdict(set)

for marca in todas_marcas:
    norm = normalize(marca)
    for tok in norm.split():
        if (len(tok) >= MIN_LEN
                and tok not in JA_FILTRADOS
                and not tok.isdigit()
                and not re.match(r'^[A-Z]{1,2}$', tok)):   # siglas de 1-2 letras
            token_marcas[tok].add(marca)

# Filtrar por frequência mínima
candidatos = {
    tok: marcas
    for tok, marcas in token_marcas.items()
    if len(marcas) >= MIN_FREQ
}
print(f'  {len(candidatos)} tokens com freq >= {MIN_FREQ}')

# Montar DataFrame
rows = []
for tok, marcas in sorted(candidatos.items(), key=lambda x: -len(x[1])):
    exemplos = sorted(marcas)[:5]
    rows.append({
        'TERMO':          tok,
        'FREQUENCIA':     len(marcas),
        'COMUM':          'SIM' if tok in COMUNS_USUARIO else '',
        'EXEMPLOS_MARCAS': ' | '.join(exemplos),
        'TODAS_AS_MARCAS': ' | '.join(sorted(marcas)),
    })

df_out = pd.DataFrame(rows)
print(f'\nTop 30 tokens mais frequentes:')
print(df_out[['TERMO','FREQUENCIA','COMUM']].head(30).to_string(index=False))

OUTPUT = './relatorio_v2/Termos_Candidatos_Comuns.xlsx'
print(f'\nSalvando {OUTPUT} ...')

with pd.ExcelWriter(OUTPUT, engine='openpyxl') as writer:
    df_out[['TERMO','FREQUENCIA','COMUM','EXEMPLOS_MARCAS']].to_excel(
        writer, index=False, sheet_name='Termos'
    )
    df_out[['TERMO','FREQUENCIA','COMUM','TODAS_AS_MARCAS']].to_excel(
        writer, index=False, sheet_name='Termos_Completo'
    )

    from openpyxl.styles import PatternFill, Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    for sheet_name in ['Termos', 'Termos_Completo']:
        ws = writer.sheets[sheet_name]

        # Header
        for cell in ws[1]:
            cell.fill  = PatternFill('solid', fgColor='2F4F8F')
            cell.font  = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center')

        # Colorir linhas pré-marcadas como COMUM
        amarelo = PatternFill('solid', fgColor='FFF2CC')
        cinza   = PatternFill('solid', fgColor='F2F2F2')

        col_comum = next(
            (i+1 for i, c in enumerate(ws[1]) if c.value == 'COMUM'), 3
        )

        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            comum_val = ws.cell(i, col_comum).value or ''
            fill = amarelo if comum_val == 'SIM' else (cinza if i % 2 == 0 else None)
            for cell in row:
                if fill:
                    cell.fill = fill
                cell.alignment = Alignment(wrap_text=True, vertical='top')

        # Larguras
        ws.column_dimensions['A'].width = 22   # TERMO
        ws.column_dimensions['B'].width = 12   # FREQUENCIA
        ws.column_dimensions['C'].width = 10   # COMUM
        ws.column_dimensions['D'].width = 80   # EXEMPLOS

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:D{ws.max_row}'

        # Instrução na linha 1 da coluna COMUM
        ws.cell(1, col_comum).value = 'COMUM\n(edite: SIM/NAO)'

print('Concluído.')
