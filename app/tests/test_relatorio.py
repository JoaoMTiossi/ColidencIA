"""Testes do agrupamento de vigilância de termo desgastado na aba Relatório."""
from __future__ import annotations

import os

import openpyxl

from app.pipeline.relatorio import gerar_xlsx


def _alerta(
    marca_base: str,
    marca_rpi: str,
    classificacao: str,
    nivel: str,
    ncl: int = 35,
) -> dict:
    return {
        "processo_base": "900000001",
        "marca_base": marca_base,
        "ncl_base": ncl,
        "ncl_versao_base": 12,
        "titular_base": "Cliente Ltda",
        "processo_rpi": "900000002",
        "marca_rpi": marca_rpi,
        "ncl_rpi": ncl,
        "despacho_codigo": "IPAS009",
        "despacho_nome": "Publicação para oposição (exame formal concluído)",
        "tipo_acao": "OPOSICAO",
        "classificacao": classificacao,
        "nivel": nivel,
        "score_final": 0.55,
    }


def _stats() -> dict:
    return {
        "rpi_numero": "2900",
        "rpi_data": "01/07/2026",
        "total_carteira": 100,
        "total_rpi": 200,
    }


def test_agrupamento_vigilancia_reduz_linhas_na_aba_cliente(tmp_path):
    resultados = [
        _alerta("MEGA STORE", "MEGA CASA IMPORTS", "BAIXA", "VIGIAR"),
        _alerta("MEGA STORE", "MEGA TOLDOS", "BAIXA", "VIGIAR"),
        _alerta("MEGA STORE", "MEGA MARMORES", "BAIXA", "VIGIAR"),
        _alerta("ACME SOFT", "ACME SOFTWARE", "ALTA", "ALTA"),
    ]

    out_path = os.path.join(tmp_path, "relatorio.xlsx")
    gerar_xlsx(resultados, _stats(), out_path)

    wb = openpyxl.load_workbook(out_path)

    ws_cliente = wb["Relatório"]
    # Cabeçalho de colunas na linha 7; dados a partir da linha 8.
    linhas_cliente = [
        row for row in ws_cliente.iter_rows(min_row=8, values_only=True)
        if any(v not in (None, "") for v in row)
    ]
    assert len(linhas_cliente) == 2

    resumo = [l for l in linhas_cliente if "vigilância" in (l[5] or "")]
    assert len(resumo) == 1
    assert "'MEGA': 3 marcas da RPI (vigilância)" == resumo[0][5]
    assert resumo[0][1] == "MEGA STORE"

    individual = [l for l in linhas_cliente if l is not resumo[0]]
    assert len(individual) == 1
    assert individual[0][1] == "ACME SOFT"
    assert individual[0][5] == "ACME SOFTWARE"

    ws_tecnica = wb["Análise Técnica"]
    linhas_tecnica = [
        row for row in ws_tecnica.iter_rows(min_row=2, values_only=True)
        if any(v not in (None, "") for v in row)
    ]
    assert len(linhas_tecnica) == 4


def test_alertas_altos_nunca_sao_agrupados(tmp_path):
    resultados = [
        _alerta("MEGA STORE", "MEGA CASA IMPORTS", "MEDIA", "MEDIA"),
        _alerta("MEGA STORE", "MEGA TOLDOS", "ALTA", "ALTA"),
    ]
    out_path = os.path.join(tmp_path, "relatorio2.xlsx")
    gerar_xlsx(resultados, _stats(), out_path)

    wb = openpyxl.load_workbook(out_path)
    ws_cliente = wb["Relatório"]
    linhas_cliente = [
        row for row in ws_cliente.iter_rows(min_row=8, values_only=True)
        if any(v not in (None, "") for v in row)
    ]
    # MEDIA e ALTA continuam linha a linha, sem agrupamento.
    assert len(linhas_cliente) == 2
    marcas_rpi = {l[5] for l in linhas_cliente}
    assert marcas_rpi == {"MEGA CASA IMPORTS", "MEGA TOLDOS"}
