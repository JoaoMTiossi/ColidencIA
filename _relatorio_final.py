"""
Gera relatório final limpo aplicando:
  - Camada D: remove BAIXO_RISCO em classes correlatas (≠) com score < 0.70
  - Promoção: 43 BAIXO_RISCO com score >= 0.75 e spec >= 0.20 → MEDIO_RISCO
  - Remove FALSO_POSITIVO do relatório operacional
  - Mantém apenas ALTO_RISCO e MEDIO_RISCO
"""
import sys, os, re
sys.path.insert(0, os.path.dirname(__file__))

import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_DIR = './relatorio_final'
os.makedirs(OUTPUT_DIR, exist_ok=True)

def get_cl(s):
    m = re.search(r'NCL\(13\)\s*(\d+)', str(s))
    return int(m.group(1)) if m else 0

def score_num(v):
    try: return float(v)
    except: return -1.0

# ---------------------------------------------------------------------------
print('Carregando análise IA v3...')
df = pd.read_excel('./relatorio_v3/Relatorio_Analisado_IA_v3_RPI2886.xlsx')
print(f'  {len(df)} casos carregados')

df['CL_CLI_N']    = df['CLASSE CLIENTE'].apply(get_cl)
df['CL_3O_N']     = df['CLASSE TERCEIRO'].apply(get_cl)
df['SCORE_SPEC_N'] = df['SCORE_SPEC'].apply(score_num)
df['MESMA_CLASSE'] = df['CL_CLI_N'] == df['CL_3O_N']

# ---------------------------------------------------------------------------
# AJUSTE 1: promover BAIXO_RISCO com score >= 0.75 e spec >= 0.20
# ---------------------------------------------------------------------------
mask_promover = (
    (df['VEREDITO'] == 'BAIXO_RISCO') &
    (df['SCORE_NOME'] >= 0.75) &
    (df['SCORE_SPEC_N'] >= 0.20)
)
n_promovidos = mask_promover.sum()
df.loc[mask_promover, 'VEREDITO']   = 'MEDIO_RISCO'
df.loc[mask_promover, 'PRIORIDADE'] = 2
df.loc[mask_promover, 'MOTIVO_IA']  = df.loc[mask_promover, 'MOTIVO_IA'].apply(
    lambda m: '[PROMOVIDO: score alto + spec compatível] ' + m
)
print(f'  Promovidos BAIXO → MEDIO: {n_promovidos}')

# ---------------------------------------------------------------------------
# AJUSTE 2: descartar Camada D (BAIXO_RISCO, classes correlatas ≠, score < 0.70)
# ---------------------------------------------------------------------------
mask_camada_d = (
    (df['VEREDITO'] == 'BAIXO_RISCO') &
    (~df['MESMA_CLASSE']) &
    (df['SCORE_NOME'] < 0.70)
)
n_camada_d = mask_camada_d.sum()
print(f'  Camada D descartada: {n_camada_d}')

# ---------------------------------------------------------------------------
# AJUSTE 3: remover FALSO_POSITIVO e BAIXO_RISCO restante do relatório operacional
# ---------------------------------------------------------------------------
df_relevante = df[df['VEREDITO'].isin(['ALTO_RISCO', 'MEDIO_RISCO'])].copy()
df_relevante = df_relevante.sort_values(['PRIORIDADE', 'SCORE_NOME'], ascending=[True, False])
df_relevante = df_relevante.drop(columns=['CL_CLI_N','CL_3O_N','SCORE_SPEC_N','MESMA_CLASSE'])

# ---------------------------------------------------------------------------
# Estatísticas finais
# ---------------------------------------------------------------------------
contagem = df_relevante['VEREDITO'].value_counts()
total    = len(df_relevante)

print()
print('='*65)
print('RELATÓRIO FINAL — CASOS ENVIADOS AO ADVOGADO')
print('='*65)
print(f'  ALTO_RISCO  : {contagem.get("ALTO_RISCO",0):>4}  → avaliar oposição')
print(f'  MEDIO_RISCO : {contagem.get("MEDIO_RISCO",0):>4}  → monitoramento ativo')
print(f'  TOTAL       : {total:>4}')
print()

print('  DISTRIBUIÇÃO POR CLASSE (ALTO+MÉDIO):')
por_classe = df_relevante.groupby(df_relevante['CLASSE CLIENTE'].str.extract(r'(\d+)$')[0]).size()
por_classe = por_classe.sort_values(ascending=False)
labels = {
    '35':'serviços comerciais','41':'educação/entret.','36':'seguros/financeiro',
    '37':'construção','44':'saúde/veterinária','42':'TI/científico',
    '43':'alimentação','39':'transporte','30':'alimentos/confeit.','38':'telecom',
}
for cl, cnt in por_classe.head(10).items():
    print(f'    cl.{cl:>2} ({labels.get(cl,""):22}): {cnt:>4}')

print()
print('  TOP 15 ALTO_RISCO:')
for _, r in df_relevante[df_relevante['VEREDITO']=='ALTO_RISCO'].head(15).iterrows():
    print(f"    [{r['SCORE_NOME']:.3f} spec={r['SCORE_SPEC']}] "
          f"cl{r['CLASSE CLIENTE'][-2:]}  "
          f"{r['MARCA CLIENTE']!r:38} x {r['MARCA TERCEIRO']!r}")

# ---------------------------------------------------------------------------
# Salvar Excel final (uma aba por prioridade + aba completa)
# ---------------------------------------------------------------------------
OUTPUT = f'{OUTPUT_DIR}/Relatorio_Final_IA_RPI2886.xlsx'
print(f'\nSalvando {OUTPUT} ...')

cores = {'ALTO_RISCO': 'FF9999', 'MEDIO_RISCO': 'FFE599'}

with pd.ExcelWriter(OUTPUT, engine='openpyxl') as writer:

    # Aba COMPLETA
    df_relevante.to_excel(writer, index=False, sheet_name='Todos')

    # Aba só ALTO_RISCO
    df_relevante[df_relevante['VEREDITO']=='ALTO_RISCO'].to_excel(
        writer, index=False, sheet_name='Alto_Risco'
    )

    # Aba só MEDIO_RISCO
    df_relevante[df_relevante['VEREDITO']=='MEDIO_RISCO'].to_excel(
        writer, index=False, sheet_name='Medio_Risco'
    )

    # Abas por classe (top 5 classes com mais casos)
    top_classes = por_classe.head(5).index.tolist()
    for cl in top_classes:
        sub = df_relevante[
            df_relevante['CLASSE CLIENTE'].str.contains(f'NCL\\(13\\) {cl}\\b', regex=True)
        ]
        if len(sub):
            nome_aba = f'Cl{cl}_{labels.get(cl,"")[:10]}'.replace("/","").replace("\\","").replace("?","").replace("*","").replace("[","").replace("]","").replace(":","")[:31]
            sub.to_excel(writer, index=False, sheet_name=nome_aba)

    # Formatar todas as abas
    widths = {
        'PROCESSO CLIENTE':16,'MARCA CLIENTE':38,'CLASSE CLIENTE':14,
        'TITULAR CLIENTE':34,'PROCESSO TERCEIRO':16,'MARCA TERCEIRO':38,
        'CLASSE TERCEIRO':14,'VEREDITO':16,'PRIORIDADE':10,
        'SCORE_NOME':11,'SCORE_NUCLEO':12,'SCORE_SPEC':11,'TOKEN_DIST':20,
        'ESPEC_CLIENTE':46,'ESPEC_TERCEIRO':46,'MOTIVO_IA':85,
    }

    hf   = PatternFill('solid', fgColor='1F3864')
    hfont = Font(color='FFFFFF', bold=True)

    for sheet_name in writer.sheets:
        ws = writer.sheets[sheet_name]
        for cell in ws[1]:
            cell.fill = hf
            cell.font = hfont
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

        try:
            col_v = next(c.column for c in ws[1] if c.value == 'VEREDITO')
        except StopIteration:
            col_v = None

        for row in ws.iter_rows(min_row=2):
            v = row[col_v-1].value if col_v else ''
            fill = PatternFill('solid', fgColor=cores.get(v or '', 'FFFFFF'))
            for cell in row:
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            if col_v:
                row[col_v-1].font = Font(bold=True)

        for col in ws.iter_cols(1, ws.max_column):
            h = col[0].value or ''
            ws.column_dimensions[col[0].column_letter].width = widths.get(h, 14)

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

print('Concluído.')
print()
print('='*65)
print(f'ARQUIVO: {OUTPUT}')
print(f'ABAS: Todos | Alto_Risco | Medio_Risco | + top 5 classes')
print('='*65)
